"""Formal six-chapter v4 report; embeds the complete visualization manifest."""
from pathlib import Path
import csv,json,hashlib,zipfile
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
ASSETS=ROOT/'output/report_visualizations/20261006_v4'
OUT=ROOT/'output/docx/KAMP_스토리_v4_시각화반영_20261006.docx'
BASE=Path(r'C:/Users/Playdata/Desktop/프로젝트/제조AI 경진대회/26KAMP_결과보고서_초안보완_v2_20261006.txt')
FEEDBACK=Path(r'C:/Users/Playdata/Desktop/프로젝트/제조AI 경진대회/26KAMP_초안v2_피드백_시각화방향_20261006.pdf')
manifest=json.loads((ASSETS/'manifest.json').read_text(encoding='utf-8'))
audit=json.loads((ASSETS/'notebook_execution_audit.json').read_text(encoding='utf-8'))
figures={f['id']:f for f in manifest['figures']};used=[];transcript=[]
def rows(path):
    with path.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
metrics=rows(ASSETS/'frozen_model_metrics.csv');policy=rows(ASSETS/'rg3_actual_policy_comparison.csv')

d=Document();s=d.sections[0]
s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(1.5);s.bottom_margin=Cm(1.5)
s.left_margin=Cm(2);s.right_margin=Cm(2);s.header_distance=Cm(1);s.footer_distance=Cm(1)
for name,size,bold in [('Normal',14,False),('Title',23,True),('Subtitle',14,False),('Heading 1',16,True),('Heading 2',14,True),('Heading 3',13,True),('Caption',10,False)]:
    style=d.styles[name];style.font.name='휴먼명조';style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'휴먼명조')
    style.paragraph_format.line_spacing=1.6 if name=='Normal' else 1.15
    style.paragraph_format.space_after=Pt(6)
    if name.startswith('Heading'):
        style.paragraph_format.space_before=Pt(12);style.paragraph_format.keep_with_next=True
    for p in list(style.element.get_or_add_pPr()):
        if p.tag in (qn('w:pBdr'),qn('w:shd')):style.element.get_or_add_pPr().remove(p)
normal=d.styles['Normal'];normal.paragraph_format.widow_control=True
header=s.header.paragraphs[0];header.text='26KAMP  CN7 RG3 결과보고서 초안 v4';header.style='Caption'
footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def p(text,style=None):
    paragraph=d.add_paragraph(text,style=style);transcript.append(text);return paragraph
def h(text,level=2,new_page=False):
    para=d.add_heading(text,level=level);para.paragraph_format.page_break_before=new_page
    transcript.extend(['',text,'']);return para
def chapter(number,title,points):
    h(f'제{number}장. {title} [{points}점]',1,new_page=True)
def table(headers,data,widths=None):
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=widths or [17/len(headers)]*len(headers)
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for ri,values in enumerate([headers]+data):
        row=t.rows[0] if ri==0 else t.add_row();prop=row._tr.get_or_add_trPr();prop.append(OxmlElement('w:cantSplit'))
        if ri==0:prop.append(OxmlElement('w:tblHeader'))
        for cell,text,w in zip(row.cells,values,widths):
            cell.width=Cm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cell.text=str(text)
            cp=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for edge,value in [('top',100),('bottom',100),('left',90),('right',90)]:
                item=OxmlElement('w:'+edge);item.set(qn('w:w'),str(value));item.set(qn('w:type'),'dxa');m.append(item)
            cp.append(m);borders=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                item=OxmlElement('w:'+edge);item.set(qn('w:val'),'single');item.set(qn('w:sz'),'4');item.set(qn('w:color'),'BFC5CC');borders.append(item)
            cp.append(borders)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E9EDF1');cp.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.line_spacing=1.1;para.paragraph_format.space_after=Pt(0)
                for run in para.runs:run.font.size=Pt(10.5);run.bold=ri==0
        transcript.append(' | '.join(map(str,values)))
    p('').paragraph_format.space_after=Pt(0)
    return t
def figure(key,section=None):
    assert key not in used;f=figures[key];used.append(key);number=len(used)
    h(section or f'그림 {number} {f["title"]}',2,new_page=True)
    para=p(f'그림 {number}. {f["title"]}','Caption');para.paragraph_format.keep_with_next=True
    picture=p('');picture.paragraph_format.line_spacing=1;picture.paragraph_format.keep_with_next=True
    picture.add_run().add_picture(str(ROOT/f['path']),width=Cm(17))
    context_text=f['context']
    if key in ('cn7_error_conditions','rg3_error_conditions'):
        product='CN7 전체 위험3+FP8' if key.startswith('cn7') else 'RG3 위험5+상위점수 FP10'
        context_text=context_text.split('/ CN7 전체')[0]+' / '+product+' / 표시만 IQR 변환'
    if key in ('ocsvm_confusion','lr_confusion'):
        kind=key.split('_')[0]
        context_text+=' / '+ ' ; '.join(r['dataset'].upper()+' '+r['version']+' t='+r['threshold'] for r in metrics if r['model']==kind)
    context=p(context_text,'Caption');context.paragraph_format.keep_with_next=True
    notebook=next(k for k,g in manifest['notebooks'].items() if g==f['group'])
    origin=p('산출 노트북: '+notebook,'Caption');origin.paragraph_format.keep_with_next=True
    p(f['interpretation'])
    transcript.extend([f'[그림 {number}: {key}]',f['context'],f['interpretation']])

p('사출성형 공정의 품질 위험 분석과 검사 지원', 'Title')
p('제6회 K 인공지능 제조데이터 분석 경진대회\n일반국민 대학 원 생 부문 결과보고서 초안 v4', 'Subtitle')
p('작성일 2026년 10월 6일')
p('CN7과 RG3의 공정 데이터에서 불량 관측 이력이 있는 입력 패턴을 분석하고, 예측 점수를 검사 대상 선정과 추가 정보 확보에 연결하였다. 동일 입력의 상충 라벨을 보존하고 패턴 단위로 학습·평가를 분리하였다. IF, OCSVM, LR, RF와 관계 잔차 및 스태킹 실험을 제품별로 비교하였다.')
p('CN7의 고정 RF는 기존 Test 위험 패턴 3개를 모두 탐지하였으나 정상 패턴 8개를 오탐하였다. RG3의 OCSVM·LR은 위험 5개를 탐지하는 대신 정상 오탐이 많았다. 실제 우선검사·무작위 검사 규칙의 후향 비교에서도 개선 근거가 확보되지 않아 RG3의 최종 검사 정책은 확정하지 않았다.')
p('원본 관측치와 입력 패턴의 평가 단위를 구분하였다. 고정 Test는 기존에 관찰한 자료이므로 모든 Test 결과는 후속 평가로 한정하였다. 독립 현장 자료의 재학습 전후 성능과 비용 절감은 검증되지 않았다.')
table(['제출 정보','기입 상태'],[['팀명','미기입'],['대표자 및 팀원','미기입'],['서명 및 필수 설문 완료 화면','최종 제출 자료에 별도 첨부']], [5,12])

h('목차',1,new_page=True)
chapters=[('데이터 이해 및 진단',15),('AI 예측모델 개발 및 성능평가',40),('영향요인 및 오류분석',15),('현장 활용방안',10),('창의성 및 차별성',10),('코드 구성 및 재현성',10)]
for i,(title,point) in enumerate(chapters,1):p(f'제{i}장. {title} [{point}점]')
p('부록 A 전체 시각화 결과\n부록 B 제출 전 미완료 사항과 출처')
p('본문에는 판단의 근거가 되는 핵심 그림을 배치하였다. 각 노트북의 추가 그림은 부록 A에 전부 수록하였다. 그림별로 자료, 평가 단위, 표본 수, 모델 및 판정 조건을 기록하였다.')

chapter(1,*chapters[0]);h('1.1 과제와 분석 대상')
p('사출성형 공정의 품질 이상을 최종검사 이전에 식별하고, 모델이 작동하거나 실패하는 공정조건을 분석하는 것을 목적으로 하였다. 시간·위치·속도·압력·온도 관련 24개 입력을 사용하였다. PassOrFail은 0이 양품, 1이 불량이며 품질 라벨과 원본 식별자는 예측 입력에서 제외하였다.')
table(['구분','원본 행','양품','불량','고유 패턴','위험 패턴'],[['CN7','1211','1194','17','606','14'],['RG3','1182','1157','25','591','25']],[2,3,3,3,3,3])
p('비라벨 자료는 CN7 35239행, RG3 35941행이다. 동일 입력을 통합한 비라벨 패턴은 각각 28771개와 29707개였다. 품질 라벨이 없어 정상 자료로 간주하지 않았다. 원본 불량 비율은 CN7 1.40%, RG3 2.12%로, 전체 정확도만으로 위험 탐지 성능을 판단하기 어렵다.')
h('1.2 상충 라벨과 예측 대상')
p('CN7의 위험 패턴 14개 중 11개와 RG3의 위험 패턴 25개 모두에서 동일 입력의 양품·불량 라벨이 상충하였다. CN7에는 불량 전용 패턴 3개가 별도로 존재하였다. 같은 입력을 사용하는 결정적 분류기는 상충 패턴의 개별 원본 행을 구별할 수 없다. 생산 맥락이 부족한 상태에서 상충을 라벨 오류로 단정하지 않았다.')
p('동일한 24개 입력을 하나의 패턴으로 집계하고 max(label)로 위험 이력을 정의하였다. 이때 지도학습의 대상은 개별 제품의 불량 여부가 아니라 불량 관측 이력이 있는 입력 패턴이다. 원본 연결 키와 정상·불량 관측 수를 유지하여 관측치 검사 평가를 별도로 수행하였다.')
h('1.3 변수 분포와 검증 설계')
p('라벨 입력의 결측·무한값과 정확한 패턴 중복을 확인하였다. CN7에서는 위치·배압·금형온도 관련 변수의 라벨별 분포 차이가 관측되었다. RG3에서는 일부 시간 변수의 라벨 자료 고유 값 수가 적고 비라벨 자료의 다양성이 더 컸다. 상관된 변수의 중복된 정보는 거리 및 관계식 분석에서 별도로 고려하였다.')
p('패턴 단위의 고정 Test 20%와 개발 자료 4-fold를 적용하였다. CN7 개발 자료는 정상 473·위험 11개, Test는 정상 119·위험 3개였다. RG3 개발 자료는 정상 452·위험 20개, Test는 정상 114·위험 5개였다. 후보 선택은 개발 자료에서 수행하였으며 동일 패턴의 교차 포함을 차단하였다.')
h('1.4 해석 가능한 범위')
p('입력은 이미 표준화된 좌표로 제공되었으나 최초 표준화 방법과 기준집단은 확인되지 않았다. 라벨·비라벨의 기준 정합도 확인되지 않았다. 따라서 좌표 차이를 물리적 온도·압력 차이로 환산하거나 비라벨 분포 차이를 현장 품질 악화로 확정하지 않았다.')
p('실제 생산시각, LOT, 설비 및 제품 추적 정보가 충분하지 않아 원본 행 순서나 Unnamed: 0을 생산시각·현장 제품 ID로 해석하지 않았다. 전체 자료의 EDA 결과는 탐색 근거이며 독립 Test 성능 또는 인과적 불량 원인이 아니다.')
figure('eda_counts');figure('preprocess_conflicts');figure('split_counts')

chapter(2,*chapters[1]);h('2.1 모델과 선택 기준')
table(['모델','학습과 특징','선택 또는 판정'],[
 ['IF','개발 정상 / 300 trees / 분산 제거·정상 기준 변환','불량 임계값 없음 / 이상 순위'],
 ['OCSVM','개발 정상 / 7개 특징 구성 / nu·gamma 탐색','개발 4-fold OOF F1 / 숫자 임계값 고정'],
 ['LR','정상·위험 / 8개 구성 / C·클래스 가중치 탐색','개발 4-fold OOF F1 / 숫자 임계값 고정'],
 ['고정 RF','CN7 사전 기준 / 300 trees / depth5 / leaf5 / balanced','전체 탐색 선택 모델 아님 / t=0.5'],
 ['연구 RF','7개 구성 / 구성당108설정 / 1001 임계값','v1 기존 탐색 결과 / 수정 후 v2 전체 재평가 미완료']], [2.8,7.3,6.9])
p('OCSVM의 기본 학습은 IF와 독립적이다. if_desc 비교 모드에서만 IF가 생성한 정상 학습 순서를 사용한다. IF는 공통 추론·분포 감지·재학습·평가와 스태킹 경로에도 포함된다. LR과 RF 기본 학습에서 IF 점수를 필수 입력으로 사용하지 않는다.')
p('OCSVM은 7개 구성과 nu 28개·gamma 22개를 비교하며, LR은 8개 구성과 C 15개·가중치 2개를 비교하였다. 선택 OOF에는 후보 탐색의 선택 편향이 포함된다. Test 결과를 보고 설정·임계값을 다시 선택하지 않았다.')
h('2.2 고정 Test의 분류 결과')
table(['제품 모델','임계값','TP','FP','FN','TN','F1'],[
 [r['dataset'].upper()+' '+r['model'].upper(),f'{float(r["threshold"]):.6g}',*[str(int(float(r[k]))) for k in ['TP','FP','FN','TN']],f'{float(r["F1"]):.4f}'] for r in metrics if r['threshold']],[3.7,3.0,1.7,1.7,1.7,1.7,1.5])
p('CN7 고정 RF의 재현율은 100%, 정밀도는 27.27%, 정상 오탐률은 6.72%였다. LR·OCSVM은 기존 임계값에서 위험 3개를 모두 놓쳤다. RG3 OCSVM은 위험 5개를 탐지하였으나 정상 89개를 오탐하였으며, LR의 정상 오탐은 110개였다. RG3의 재현율만으로 적용 가능성을 판단할 수 없었다.')
p('IF는 불량 분류 임계값이 없는 이상 순위 모델이므로 임의의 임계값을 만들어 혼동행렬을 작성하지 않았다. F1 분류와 검사 예산에 따른 순위 평가는 구분하였다. 위험 한 개의 탐지 여부가 Test 재현율을 CN7 약 33.3%p, RG3 20%p 바꾼다.')
h('2.3 모델 버전과 판정 조건')
table(['제품','모델','저장 버전'],[[r['dataset'].upper(),r['model'].upper(),r['version']] for r in metrics],[2,2,13])
p('CN7 OCSVM은 legacy_without_cycle, nu=0.001, gamma 배수=0.0385388을 사용하였다. RG3 OCSVM은 group_pca, nu=0.07, gamma 배수=10을 사용하였다. LR은 두 제품 모두 thermal 구성이며 CN7은 C=10·t=0.921, RG3는 C=0.316228·t=0.012이고 클래스 가중치는 적용하지 않았다.')
h('2.4 RF 표준화 수정과 실험 범위')
p('RF 연구 코드에는 제공된 표준화 좌표를 다시 표준화하지 않도록 수정한 v2가 반영되어 있다. PCA는 제공 좌표를 중심화하여 적합하며, 별도의 재표준화와 구분된다. 이번 시각화 추가에서는 RF 학습·변환 코드를 변경하지 않았다. 원본 제공 좌표의 최초 표준화가 한 번 수행되었다는 사실과 그 기준은 자료 제공 측 확인이 필요하다.')
p('기존 manual_seven_scenarios_v1의 OOF·Test·중요도는 수정 이전 결과이다. CN7 v1은 pressure·t=0.222에서 개발 F1 0.5333, Test F1 0이었다. RG3 v1은 group_pca·t=0.47에서 개발 F1 0.1875, Test F1 0이었다. 이 수치를 표준화 수정 후의 성능으로 옮기지 않았다. 현재 고정 CN7 RF와 연구 RF도 다른 모델이다.')
h('2.5 관계 잔차와 스태킹')
p('관계 잔차 실험에서는 정상 자료의 한 변수를 다른 입력으로 예측하고, 예상값과 실제값의 이탈을 특징으로 사용하였다. 정상 관계 적합용 표본과 분류 학습 표본을 분리하고 GroupKFold로 관계의 예측력을 확인하였다. CN7의 검증 선택 ridge_residual_lr은 Test F1 0, RG3의 extra_residual_lr은 TP1·FP31·FN4·F1 0.0541이었다.')
p('스태킹은 RF·관계 잔차·IF의 점수를 외부4·내부3 fold의 nested 구조로 결합하였다. 개발 상위 10% 위험 발견 수를 주 선택 지표로 사용하였다. CN7 선택 스태킹의 Test는 TP1·FP0·FN2·TN119·F1 0.5였고, RG3는 TP0·FP13·FN5·TN101·F1 0이었다. 같은 실험 안의 개별 모델과 비교하였으며 고정 운영 모델과 합치지 않았다.')
h('2.6 제품별 잠정 적용 판단')
p('CN7 고정 RF는 검사 우선순위의 잠정 기준으로 사용할 근거가 있으나, 개별 제품의 확정 판정과 최종 승격은 독립 검사 자료의 검증이 남아 있다. RG3는 높은 오탐과 약한 점수 구분력 때문에 최종 분류 모델·검사 정책을 확정하지 않았다. 수정 후 연구 RF의 전체 탐색과 최종 평가도 미완료 범위이다.')
figure('lr_scores');figure('ocsvm_confusion')

chapter(3,*chapters[2]);h('3.1 실제 성공과 실패 사례')
p('CN7 Test 위험 패턴은 50·51·57번이며 세 패턴 모두 상충 라벨 이력이 있다. 고정 RF는 세 패턴을 탐지하였으나 정상 14·28·35·42·66·72·75·76번을 오탐하였다. LR·OCSVM은 같은 위험 세 패턴을 기존 임계값에서 놓쳤다. Test 결과를 공정값과 연결하여 성공·실패 조건을 대조하였다.')
cases=rows(ASSETS/'cn7_actual_error_cases.csv');selected=[r for r in cases if r['category']!='TN']
table(['패턴','판정','RF 점수','배압','금형온도3','금형온도4'],[[r['pattern_row'],r['category'],f'{float(r["score"]):.4f}',f'{float(r["Max_Back_Pressure"]):.3f}',f'{float(r["Mold_Temperature_3"]):.3f}',f'{float(r["Mold_Temperature_4"]):.3f}'] for r in sorted(selected,key=lambda r:(-int(r['label']),int(r['pattern_row'])))],[2,2,3,3,3.5,3.5])
p('표의 값은 제공된 표준화 좌표이다. 그림의 중앙값·IQR 변환은 사례 비교를 위한 표시 목적이며 모델 입력에 재표준화를 적용한 것이 아니다. 중요도·계수와 관측 조건은 연관 근거로 한정하였다.')
h('3.2 개발 정상 구간에 따른 오류 건수')
intervals=rows(ASSETS/'error_interval_counts.csv')
table(['제품','배압 구간','Test 수','위험 수','FP','FN'],[[r['dataset'].upper(),r['interval'],r['N'],r['risk'],r['FP'],r['FN']] for r in intervals],[2,6,2.3,2.3,2.2,2.2])
p('구간 경계는 개발 정상 자료의 배압 사분위수로 고정하였다. CN7의 오탐 8개는 가장 낮은 배압 구간 38개 안에 모두 존재하였다. 해당 구간에는 위험 2개도 함께 있어 낮은 배압 하나만으로 불량 여부를 확정할 수 없다. 나머지 위험 한 개는 다음 배압 구간에 위치하였다.')
p('RG3 OCSVM의 정상 오탐은 네 배압 구간에 각각 28·23·20·18개로 분포하였다. 위험과 정상 예측 점수가 중첩되어 하나의 국소 구간 오류만으로 설명하기 어려웠다. 낮은 임계값에 따른 오탐 증가와 점수 순위의 구분력 부족을 함께 확인하였다.')
h('3.3 공정변수의 결합 조건')
p('배압과 금형온도, 사출압력과 쿠션위치의 두 변수 분포에 TP·TN·FN·FP를 함께 표시하였다. CN7 위험 세 패턴과 오탐 정상 패턴은 금형온도 및 저배압 영역에서 겹쳤다. RG3 위험 패턴과 다수의 오탐도 넓은 입력 영역에 함께 분포하였다. 관측된 영역을 새 판정 규칙으로 선택하지 않았으며 물리적 불량 원인으로 확정하지 않았다.')
joint=rows(ASSETS/'error_joint_counts.csv')
table(['제품','배압 금형온도3','표본','위험','TP','FP','FN','TN'],[[r['dataset'].upper(),r['back_pressure']+' / '+r['mold_temperature'],r['N'],r['risk'],r['TP'],r['FP'],r['FN'],r['TN']] for r in joint],[2,4.2,1.8,1.8,1.8,1.8,1.8,1.8])
cn7j=next(r for r in joint if r['dataset']=='cn7');rg3j=next(r for r in joint if r['dataset']=='rg3')
p(f'하위·상위는 개발 정상 중앙값으로 구분하였다. CN7 경계는 배압 {float(cn7j["pressure_cutoff"]):.4f}·금형온도3 {float(cn7j["temperature_cutoff"]):.4f}, RG3 경계는 배압 {float(rg3j["pressure_cutoff"]):.4f}·금형온도3 {float(rg3j["temperature_cutoff"]):.4f}이다. CN7의 하위 배압·상위 금형온도3 조합은 50개 패턴에 위험 3개와 FP 8개, TN 39개가 함께 존재하였다. 표본 수와 오류 건수를 함께 표시하여 희소 조건의 과도한 해석을 제한하였다.')
h('3.4 영향요인과 추가 가설의 해석')
p('LR 계수는 모델의 변환 좌표에서 다른 입력을 포함한 조건부 관계이며, RF 불순도 중요도는 해당 모델의 분기 사용 정도이다. 상관된 입력 간 중요도의 분산, 표본 수와 선택 편향을 고려하였다. 연구 RF v1의 중요도를 수정 후 v2의 근거로 사용하지 않았다.')
p('관계식의 정상 예측력과 위험 판별력은 별도로 판단하였다. RG3의 행 순서 기반 관계 잔차 결과는 생산시각이 확인되지 않은 가정 실험이다. CN7의 유사 시간 분할은 검증 위험 표본이 없어 품질 성능 비교가 성립하지 않았다. 자기학습 초안도 과거 참조 모델 경로가 없어 실행 완료나 개선 성과로 기술하지 않았다.')
figure('cn7_error_conditions');figure('cn7_error_joint');figure('rg3_error_conditions');figure('rg3_error_joint')

chapter(4,*chapters[3]);h('4.1 검사 예산과 품질 판정')
p('점수 순위는 한정된 검사량에 더 많은 위험 이력 패턴 또는 불량 관측치를 포함하는 목적으로 평가하였다. 검사량은 ceil(N×비율)로 정하고, 동점은 입력 fingerprint와 record_id 순서에 따라 처리하였다. 대상 선정에는 평가 라벨을 사용하지 않았다.')
table(['제품 모델','Test 단위','10% 검사 수','발견 위험','발견률'],[[r['dataset'].upper()+' '+r['model'].upper(),'고유 패턴',r['top10_k'],r['top10_TP'],f'{int(r["top10_TP"])/int(r["risk"]):.1%}'] for r in metrics],[4,3,3,3,4])
p('CN7 고정 RF는 13개 패턴을 검사할 때 위험 3개를 포함하였고, LR은 2개, IF·OCSVM은 0개였다. RG3 IF는 12개 패턴에 위험 2개를 포함하였고 LR·OCSVM은 0개였다. 동일 예산의 후향 결과이며 실제 현장 제품의 불량 감소 효과와 구분하였다.')
h('4.2 RG3 실제 정책 규칙의 후향 비교')
p('현행 코드의 RG3 정책은 참조 범위 이탈과 정보 부족을 우선하고, 동일 우선순위에서는 ID 순서로 대상 일부를 정한다. 나머지는 잔여 집합에서 무작위 추출한다. IF 점수 순위와 동일한 정책이 아니다. 개발 패턴에서 안내 근거를 생성한 뒤 고정 Test 대응 원본 238행·불량 5행에 실제 정책 함수를 적용하였다. 출력은 격리 폴더에 저장하였다.')
p('전체 검사 비율은 5·10·20%, 무작위 비중은 0.5로 사전에 정한 예시이다. 정책 우선 집합은 고정하고 무작위 잔여 추출을 1000회 반복하였다. 분석용 ID는 source_row 연결 키이며 실제 제품 ID로 확인되지 않았다. 현행 운영 상태의 승인 정책을 변경하지 않았다.')
table(['검사 비율','검사 수','우선+무작위 기대 발견','IF 발견','무작위 기대 발견'],[[f'{float(r["fraction"]):.0%}',r['k'],f'{float(r["policy_mean_TP"]):.3f}',r['IF_TP'],f'{float(r["random_mean_TP"]):.3f}'] for r in policy],[2.7,2.3,5.2,2.8,4])
r=policy[1];p(f'10%의 24행 검사에서 우선검사·무작위 조합의 기대 발견 수는 {float(r["policy_mean_TP"]):.3f}행, 반복 평균은 {float(r["policy_sim_mean"]):.3f}행이었다. 단순 무작위 기대값은 {float(r["random_mean_TP"]):.3f}행이고 IF 단독은 {r["IF_TP"]}행을 포함하였다. 실제 코드의 고정 시드 계획에서는 0행이 발견되었다. 해당 조합의 기대 발견률은 5.31%, 기대 검사 정밀도는 1.11%로 단순 무작위보다 개선되지 않았다.')
p('20% 예산의 한 고정 시드에서 2행을 포함하였으나 기대 발견 수는 0.561행이었다. 한 번의 유리한 추출 결과만으로 정책 효과를 주장하지 않았다. 우선순위의 입력 근거와 ID 동점 처리가 품질 위험을 충분히 구별하지 못할 수 있으므로 추가 독립 자료에서 정책을 검증해야 한다.')
h('4.3 실제 비라벨 자료의 분포 변화')
p('개발 정상 참조와 실제 비라벨 자료의 주요 입력 및 같은 IF 모델의 점수 분포를 비교하였다. 관측된 분포 차이는 라벨·비라벨의 좌표 기준 차이, 자료 구성 또는 공정 변화의 영향을 포함할 수 있다. 실제 생산시각과 품질 라벨이 없어 시간 흐름 또는 품질 악화로 확정하지 않았다.')
p('현행 분포 감지는 고정 구간의 TV, 최대 변화량 bootstrap 기준과 표본 충분성을 확인한다. 기본 기준은 원본 200행 이상·고유 패턴 40개 이상·bootstrap 200회·99% 분위수와 동일 정책·기준선에서의 연속 변화이다. 대기 상태는 정상 보증을 의미하지 않는다.')
h('4.4 재학습과 적용 판단')
p('변화 감지 이후 입력 오류·표준화 정합·정상 공정 변경 여부를 확인하고 실제 검사 라벨을 확보한 뒤 재학습 후보를 승인한다. 신규 배치의 학습 결과 개선은 독립 평가 개선과 구분한다. 기존·후보는 같은 독립 holdout, 고정 판정 규칙과 동일 검사 예산에서 FN·FP·발견 수를 비교한다.')
p('승격에는 학습·선택·기존 Test와 겹치지 않는 평가 패턴, 충분한 위험 표본 및 승인 기록이 필요하다. 정상 참조 갱신에도 실제 생산시각·공정 버전·라벨 출처와 확인된 정상 표본이 요구된다. 과거 위험 이력은 유지하며 최근 학습 범위 선택과 구분하였다. 신규 정상 라벨로 상충 위험 이력을 지우지 않는다.')
p('현재 RG3의 기준선 포인터와 단계별 검사 비율·담당자 정책은 미설정이다. CN7에는 고정 RF·IF가 활성화되어 있으나 승인 담당자 정책은 미설정이다. 미설정 조건의 보류와 기능 구현을 구분하였다. 독립 현장 배치의 기존·후보 성능 자료가 없으므로 정상 유지·개선 후 승격·개선 실패 사례를 실증 성과로 만들지 않았다.')
figure('model_budget_curves');figure('rg3_actual_policy_comparison');figure('unlabeled_distribution_change');figure('unlabeled_score_change');figure('operations_ct_decision')

chapter(5,*chapters[4]);h('5.1 분석 설계와 적용 범위')
p('상충 라벨을 보존하는 위험 이력 정의, 패턴 단위 누수 방지와 원본 관측치 연결을 하나의 흐름으로 구성하였다. CN7과 RG3의 성능 차이를 반영하여 동일한 예측·검사 정책을 강제로 적용하지 않았다. CN7은 고정 RF의 검사 순위 활용을 검토하였고, RG3는 오탐과 정보 부족을 고려하여 실제 검사 라벨 확보와 정책 검증을 우선하였다.')
h('5.2 정상 관계와 점수 결합의 기여')
p('관계 잔차는 정상 공정변수의 관계에서 이탈한 정도를 표현하였다. 스태킹은 서로 다른 점수의 보완 가능성을 nested 구조로 검토하였다. CN7 개발 상위 49개에서 위험 발견 수는 개별 RF 10개, 스태킹 11개였다. Test 상위 13개에서는 두 모델 모두 위험 3개를 포함하여 추가 발견 수가 늘지 않았다.')
p('CN7 Test의 고정 분류 규칙에서는 동일 nested 실험 RF가 TP3·FP11·F1 0.3529, 스태킹이 TP1·FP0·F1 0.5였다. 스태킹은 오탐을 줄였으나 위험 2개를 놓쳤다. 이 비교는 고정 운영 RF의 FP8·F1 0.4286과 다른 실험이다. 지표별 차이를 하나의 일관된 개선으로 표현하지 않았다.')
p('RG3 스태킹은 개발 상위 48개에서 위험 4개를 포함하였으나 Test 상위 12개에서는 0개였다. 관계 잔차의 검증 선택 모델도 위험 판별 개선이 제한적이었다. 새로운 특징이나 결합 자체를 차별성의 성능 근거로 사용하지 않고 제품·지표·분할의 범위를 기록하였다.')
h('5.3 관리 기능과 미검증 효과')
p('모델·기준선·검사 정책·라벨 출처·승인 근거를 연결하고 전환 실패 복구와 롤백 이력을 남기는 구조를 구현하였다. 로컬 승인 담당자 기록과 조건 검사는 외부 조직 인증·권한 시스템 연동과 구분하였다. 실제 라벨 오류 정정 API는 현재 구현되어 있지 않다.')
p('기능 테스트는 저장·연결·예외·승인·복구의 정확성을 확인한 근거이다. 현장 불량 감소, 검사 비용 절감과 독립 성능 개선의 실증을 대신하지 않는다. 실제 검사 수, 불량 발견 수, 처리 시간과 비용 자료가 없으므로 절감액을 산정하지 않았다.')
figure('stacking_same_experiment')

chapter(6,*chapters[5]);h('6.1 기준 코드와 진입점')
p('검토 기준은 develop_seo 브랜치의 086b02891188d1a45d474482ad2afb315458f05b이다. RF 재표준화 제거와 승인 정책 변경이 포함된 상태에서 시각화 셀과 reporting 공통 코드를 추가하였다. 기존 학습·운영 코드와 저장 모델은 유지하였다.')
table(['단계','진입점과 주요 산출물'],[
 ['EDA','01.EDA/01_eda.ipynb / EDA/cn7_defect_factor_analysis.ipynb / 26kamp.ipynb'],
 ['전처리 및 분할','02.preprocessing/02_preprocess.ipynb / 03_split.ipynb / conservative CSV·mapping·split manifest'],
 ['기본 모델','03.modeling/models/04.1_isolation_forest.ipynb / ocsvm.ipynb / 05_logistic_regression.ipynb'],
 ['RF 및 추가 실험','modeling/random_forest_*_integrated.ipynb / relationship_residual_analysis.ipynb / stacking_complementarity_experiment.py'],
 ['공통 및 운영','03.modeling/common/00_pipeline_core.ipynb / pipeline_runtime·decision_runtime·workflow_runtime·governance_runtime.py / 04_operations.ipynb'],
 ['보고서 시각화','reporting/visualizations.py / execute_notebook_sections.py / output/report_visualizations/20261006_v4/']], [3.4,13.6])
h('6.2 실행 환경과 산출물')
p('시각화 실행 환경은 Python 3.14.6, scikit-learn 1.9.0, pandas 3.0.3, NumPy 2.4.6, matplotlib 3.11.0이다. 저장 운영 모델과 같은 scikit-learn 버전으로 추론하였다. 연구 RF의 저장 CSV는 기존 환경 결과를 읽었으며 수정 후 모델을 재학습하지 않았다. 연구·운영 의존 파일의 버전이 달라 단일 제출 환경 확정이 남아 있다.')
p('시각화는 python reporting/visualizations.py로 생성한다. 각 노트북의 마지막 보고서 셀은 해당 그림과 해석을 표시한다. 전체 학습 셀을 다시 실행하지 않아도 된다. 보고서 재생성은 번들 문서 환경의 Python으로 reporting/build_report.py를 실행하고 Word 또는 호환 편집기에서 PDF로 내보내는 경로이다.')
p(f'14개 노트북에 추가한 시각화 셀을 실제 커널에서 실행하여 그림 {audit["figure_outputs"]}개의 출력을 확인하였다. 기존 모든 셀의 내용은 보존하였으며 모델·전처리 저장물 {audit["protected_artifacts_unchanged"]}개의 해시가 변경되지 않았다. 학습 셀은 재실행하지 않았다. 그림 전수와 CSV 근거·실행 기록은 별도 시각화 폴더에 보관하였다.')
h('6.3 기능 검증과 재현성 범위')
p('기존 20261006T052246Z-d5c4a94a 기록에는 143개 기능 검증 통과가 저장되어 있다. RF의 20261006T050408Z-d88dc7a0 기록에는 18개 검증 통과가 저장되어 있다. 해당 기록은 기능 동작 근거이며 이번 시각화 추가 후 전체 모델 학습을 다시 실행한 기록은 아니다.')
p('추가 시각화에서는 저장 모델의 기존 Test 지표를 대조하고, 실제 RG3 정책 함수의 격리 호출과 전 노트북의 새 셀 실행을 확인하였다. 원본 연결과 단위, 모델 버전, 표본 수 및 임계값을 대조하였다. 데이터·모델·정책의 실제 운영 포인터를 변경하지 않았다.')
p('전처리와 전체 후보 학습, 저장 모델 추론 및 예측 파일 생성의 경로는 구현되어 있다. 그러나 제출물만으로 새 환경에서 모든 후보 학습부터 최종 예측 파일까지 실행한 종단 기록은 이번 결과에 포함되지 않는다. 저장 모델의 scikit-learn 버전 검증과 요구 패키지 정합이 제출 재현의 선행 조건이다.')
h('6.4 제출 상태')
p('최종 적용 모델·검사 예산의 독립 검증, 수정 후 연구 RF의 전체 탐색, 실제 생산·검사 메타데이터 연결과 단일 제출 환경의 종단 실행은 남은 범위이다. 필수 설문 완료 화면과 대표자·팀원 서명도 제출 자료에 별도로 반영한다. 구현 완료, 설정 준비, 예측 성능 및 현장 효과를 구분하였다.')
figure('core_data_contract')

h('부록 A 전체 시각화 결과',1,new_page=True)
p('본문의 핵심 그림과 아래 추가 그림을 합하여 40개 전수를 수록하였다. 기존 모델의 불리한 결과도 유지하였다. RF 연구 v1의 결과, 정상 관계 예측력과 품질 판별 성능, 자기학습의 미실행 상태는 구분하였다.')
table(['노트북','추가 실행 그림 수'],[[r['notebook'],str(r['figure_outputs'])] for r in audit['notebooks']],[14.5,2.5])
p('상관 그림의 번호는 Clamp_Open_Position을 제외한 입력 순서이다. 1 Injection_Time, 2 Filling_Time, 3 Plasticizing_Time, 4 Cycle_Time, 5 Clamp_Close_Time, 6 Cushion_Position, 7 Plasticizing_Position, 8 Max_Injection_Speed, 9 Max_Screw_RPM, 10 Average_Screw_RPM, 11 Max_Injection_Pressure, 12 Max_Switch_Over_Pressure, 13 Max_Back_Pressure, 14 Average_Back_Pressure, 15–20 Barrel_Temperature_1–6, 21 Hopper_Temperature, 22 Mold_Temperature_3, 23 Mold_Temperature_4이다.','Caption')
remaining=[f['id'] for f in manifest['figures'] if f['id'] not in used]
for key in remaining:figure(key)

h('부록 B 제출 전 미완료 사항과 출처',1,new_page=True)
table(['항목','현재 상태'],[
 ['수정 후 연구 RF','별도 재표준화 제거 반영 / 전체 탐색·최종 성능 평가 미완료'],
 ['실제 공정조건 분석','고정 Test 사례 입력·배압 구간·결합 분포 분석 반영 / 물리 원인 확인 미완료'],
 ['RG3 실제 정책','코드 규칙의 후향 동일 예산 평가 반영 / 독립 현장 효과·최종 예산 미확정'],
 ['재학습 전후 비교','절차 구현 / 실제 독립 holdout의 기존·후보 성능 자료 없음'],
 ['생산 및 검사 정보','제품 ID·실제 시각·공정 버전·검사 라벨의 현장 연결 확인 필요'],
 ['좌표 정합','최초 표준화 방법과 라벨·비라벨 동일 기준 여부 확인 필요'],
 ['운영 준비','RG3 기준선 / 승인 담당자 / 단계별 검사 설정 미완료'],
 ['제출 재현','단일 환경 전체 학습→추론→최종 파일 종단 실행 미완료'],
 ['기존 실행 안내','일부 승인 명령에 --approval-id 누락 / 과거 OCSVM 경로와 자기학습 의존 경로 확인 필요'],
 ['라벨 정정','명시적 정정 API 미구현 / 위험 이력 보존과 별도']], [4.3,12.7])
p('설명 보완과 실제 분석 수행을 구분하였다. 기존 초안의 오류 사례·정책 비교 미완료 항목 중 수행한 후향 분석은 본문에 반영하고, 독립 현장 검증과 제출 실행은 남은 범위로 유지하였다.')
h('작성 자료와 결과 근거')
for text in [
 '기준 초안: 26KAMP_결과보고서_초안보완_v2_20261006.txt',
 '추가 피드백: 26KAMP_초안v2_피드백_시각화방향_20261006.pdf',
 '공식 목차: 경진대회 결과보고서 양식_일반국민,대학(원)생 부문.hwpx 및 과제공개 자료',
 '모델 근거: runtime/{cn7,rg3}/models/*/manifest.json 및 model.joblib',
 '실험 근거: output/logistic, output/random_forest_{cn7,rg3}/manual_seven_scenarios_v1, output/relationship_residual/20261001_relations_v1, output/stacking_complementarity/20261001_nested_v1',
 '추가 산출물: output/report_visualizations/20261006_v4/manifest.json, frozen_model_metrics.csv, error_interval_counts.csv, rg3_actual_policy_comparison.csv, notebook_execution_audit.json',
 '운영 근거: 03.modeling/common/WORKFLOW_ACCEPTANCE.md, GOVERNANCE_OPERATIONS.md 및 runtime 구현',
]:p(text,'Caption')

assert set(used)==set(figures) and len(used)==40
OUT.parent.mkdir(parents=True,exist_ok=True);d.save(OUT)
(ROOT/'output/documents').mkdir(exist_ok=True,parents=True)
(ROOT/'output/documents/KAMP_스토리_v4_시각화반영_20261006.txt').write_text('\n\n'.join(transcript),encoding='utf-8')
with zipfile.ZipFile(OUT) as z:
    media=[n for n in z.namelist() if n.startswith('word/media/')]
    assert len(media)==40
    embedded={hashlib.sha256(z.read(n)).hexdigest() for n in media}
    assert embedded=={hashlib.sha256((ROOT/f['path']).read_bytes()).hexdigest() for f in figures.values()}
report_audit=dict(docx=str(OUT),figures_embedded=len(media),figure_order=used,
                  base_sha256=hashlib.sha256(BASE.read_bytes()).hexdigest(),feedback_sha256=hashlib.sha256(FEEDBACK.read_bytes()).hexdigest())
(ASSETS/'report_audit.json').write_text(json.dumps(report_audit,ensure_ascii=False,indent=2),encoding='utf-8')
print('DOCX',OUT,'FIGURES',len(media))

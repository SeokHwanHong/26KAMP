from pathlib import Path
import json, hashlib
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output/logistic_regression_strategy'
OUT.mkdir(exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Inches(11.69), Inches(8.27)
sec.top_margin = sec.bottom_margin = Inches(.72)
sec.left_margin = sec.right_margin = Inches(.8)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    s = doc.styles[name]
    s.font.name = '맑은 고딕'
    s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.font.size = Pt(10.5 if name == 'Normal' else {'Title':22,'Subtitle':11,'Heading 1':15,'Heading 2':12}[name])
    s.paragraph_format.space_after = Pt(7)
    s.paragraph_format.line_spacing = 1.18
doc.styles['Heading 1'].paragraph_format.space_before = Pt(12)

def p(t): return doc.add_paragraph(t)
def h(t): doc.add_heading(t, 1)
def page(t): doc.add_page_break(); h(t)
def table(headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for c,v in zip(t.rows[0].cells,headers): c.text=v
    for row in rows:
        for c,v in zip(t.add_row().cells,row): c.text=str(v)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]: el.set(qn('w:'+k),v)
        borders.append(el)
    pr.append(borders)
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
        if i==0: trpr.append(OxmlElement('w:tblHeader'))
        for c,w in zip(row.cells,widths):
            c.width=Inches(w)
            cp=c._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E8EEF4' if i==0 else 'FFFFFF'); cp.append(shade)
            mar=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:w'),'90'); e.set(qn('w:type'),'dxa'); mar.append(e)
            cp.append(mar)
            for par in c.paragraphs:
                par.paragraph_format.space_after=Pt(3)
                for r in par.runs: r.font.size=Pt(9.5); r.bold=(i==0)
    p('')

stats=[]
for name in ['cn7','rg3']:
    df=pd.read_csv(ROOT/f'data/origin/moldset_labeled_{name}.csv')
    features=json.loads((ROOT/f'data/processed/{name}/preprocessing_manifest.json').read_text(encoding='utf-8'))['feature_columns']
    target=[c for c in df if c not in features and not c.startswith('Unnamed')][0]
    n1=int((df[target]==1).sum()); n0=int((df[target]==0).sum())
    stats.append((name.upper(),len(df),n0,n1,len(df[features].drop_duplicates()),f'{100*n1/len(df):.2f}%'))

doc.add_heading('로지스틱 회귀 불량 탐지 학습 전략',0)
p('CN7 및 RG3 사출성형 데이터 | 2026년 9월 28일')
p('목표는 공정 입력으로 양품과 불량품을 구분하고, 희소한 불량이 어떤 조건에서 탐지되거나 누락되는지 설명하는 것이다. 두 데이터셋은 각각 모델을 학습하며, 동일 입력 패턴을 보존한 4-fold 교차검증과 L2 규제 로지스틱 회귀를 사용한다. 불량 가중치와 판정 임계값의 영향을 분리하여 평가한다.')
h('1 데이터와 설계 결정')
table(['데이터','전체 행','정상','불량','고유 입력','불량 비율'],stats,[.8,.9,.9,.8,1.1,1.1])
p('입력은 원본 24개 공정 변수이며 라벨은 정상 0, 불량 1이다. ID, fold, 패턴 유형 및 라벨에서 계산한 정보는 입력에서 제외한다. 동일한 24개 입력값을 하나의 pattern_id로 정의하고 원본 행과 라벨을 유지한다. 고유 입력 수를 학습 행 수나 불량 수와 혼동하지 않는다.')
p('CN7의 불량 17행 중 11행은 정상과 입력이 같은 상충 패턴에 속하고, 나머지 6행은 불량 전용 3개 패턴이다. RG3의 불량 25행은 모두 정상과 같은 입력을 가진다. 동일 입력에는 동일 점수가 부여되므로, 손실 가중치만으로 이 라벨들을 분리할 수는 없다.')
h('4 fold 선택의 근거')
p('학습에 약 75%, 검증에 약 25%를 사용하여 학습 표본 확보와 검증 불량 수 사이에서 절충한다. 4가 최적이라는 주장은 하지 않는다. 그룹 제약을 무시한 단순 균등 배분으로는 CN7 검증 불량 약 4~5행, RG3 약 6~7행이 예상되지만 실제 배분은 생성 후 확인해야 한다.')
p('현재 저장된 processed 메타데이터는 두 데이터셋 모두 3-fold이다. 이 문서는 합의한 4-fold 전환을 명시한 학습 설계이며, 기존 분할 파일 변경이나 모델 학습은 아직 수행하지 않았다.')

page('2 데이터 분할과 스케일링')
p('외부 검증은 pattern_id를 그룹으로 지정한 StratifiedGroupKFold 4분할을 초기안으로 사용한다. shuffle=True와 random_state=42를 고정하고, 라벨 비율은 그룹 보존 범위에서 최대한 유지한다. 기존의 패턴 유형별 3분할 결과를 그대로 4분할로 해석하지 않는다.')
p('각 fold의 정상·불량 행 수, 고유 패턴 수, 상충 패턴 수, 불량 전용 패턴 수와 그룹 교집합을 기록한다. CN7의 불량 전용 패턴은 3개뿐이므로 4개 검증 fold 모두에 배치할 수 없다. 이를 숨기지 않고 해당 유형별 탐지 결과를 함께 제시한다.')
h('학습 구간 안에서 전처리')
p('제공 가이드에 따르면 CSV에는 파일별·컬럼별 표준화가 이미 적용되어 있다. LR 파이프라인은 학습 구간에서 상수 컬럼을 제거하고 StandardScaler를 적합한 후, 같은 변환을 해당 검증 구간에 적용한다. 학습과 검증을 각각 독립적으로 표준화하지 않는다.')
p('z_j = (x_j − 학습 평균_j) / 학습 표준편차_j')
p('기존 변환이 고정된 컬럼별 양의 선형 표준화라면, 학습 fold에서 재표준화한 값은 원자료를 해당 fold 기준으로 표준화한 값과 수학적으로 같다. 다만 제공 과정에 다른 변환이나 반올림이 있었다면 그 영향까지 되돌리지는 못한다. 비라벨 파일은 별도로 표준화되어 있으므로 추후 추론에 앞서 동일 좌표계인지 확인해야 한다.')
h('외부 평가와 내부 선택의 분리')
p('외부 4-fold의 검증 데이터는 최종 평가에만 사용한다. 각 외부 학습 구간 안에서 그룹을 보존한 내부 2-fold를 구성하여 가중치와 C를 선택한다. 내부 분할에서도 두 클래스의 존재와 패턴 누수 여부를 검증한다.')
p('기본 실험은 전체 공정 변수를 사용하고, 상수 제거 외의 특징 선택·PCA·오버샘플링은 추가하지 않는다. KS 상위 변수 등 전체 데이터 EDA에서 찾은 변수를 바로 고정하면 평가에 선택 영향이 남을 수 있으므로, 특징 선택 실험은 별도 비교로 관리한다.')

page('3 가중 로그 손실과 규제')
p('p_i = sigmoid(β₀ + βᵀz_i),  y_i ∈ {0, 1}')
p('ℓ_i = −[y_i log(p_i) + (1 − y_i) log(1 − p_i)]')
p('J = (1/n) Σ a_i ℓ_i + (λ/2) ||β||²')
p('이는 가중 로그 손실과 L2 규제를 설명하는 개념식이다. a_i는 행별 가중치이며 절편은 규제 대상에서 제외한다. 구현에서는 LogisticRegression의 C로 규제 강도를 조정하며 C가 작을수록 규제가 강하다. 구현의 정규화 관례를 확인하지 않고 λ = 1/C라고 단정하지 않는다.')
h('클래스 가중치 후보')
table(['후보','불량 대 정상 비율 r','목적'],[
('무가중','1','기준 성능 확인'),('완화 가중','√(n₀/n₁)','불량 강조를 완만하게 적용'),('균형 가중','n₀/n₁','두 클래스 총 가중치 균형')],[1.15,1.7,2.95])
p('n₀와 n₁은 매번 실제 적합에 사용하는 학습 구간에서 계산한다. 원본 전체 행 기준 비율은 CN7 약 1 / 8.38 / 70.24, RG3 약 1 / 6.80 / 46.28이지만 이 숫자를 모든 fold에 고정하지 않는다. 제곱근 가중치는 중간 비교 후보이며 최적값이라는 근거는 없다.')
p('가중치 평균을 1로 맞추려면 a₀ = n/(n₀ + r n₁), a₁ = r a₀로 설정한다. class_weight={0: a₀, 1: a₁}를 사용하고 sample_weight로 같은 보정을 중복 적용하지 않는다. 균형 가중일 때는 class_weight="balanced"와 같다.')
h('낮은 불량 비율과 학습 신호')
p('불량 비율은 초기 예측값이 아니다. 계수와 절편을 0으로 시작하면 초기 예측은 0.5이며, 로그 손실의 선형 점수에 대한 기울기는 p_i − y_i이다. 실제 불량에 0.02를 예측하면 기울기는 −0.98이므로 학습 신호가 사라지지 않는다. 가중 손실에서는 이 값에 a_i가 곱해진다.')
p('가중치는 불량의 학습 영향력을 높이지만 새로운 불량 패턴을 만들지는 않는다. focal loss와 같은 추가 재구성은 상충 라벨까지 강조할 가능성이 있으므로 첫 실험에서는 제외한다. 가중 모델의 출력은 실제 불량 발생 확률로 바로 해석하지 않고 예측 점수로 사용한다.')

page('4 모델 선택과 불량 판정')
table(['설정','초기 실험안'],[
('파이프라인','상수 제거 → StandardScaler → LogisticRegression'),
('규제와 최적화','L2 규제, lbfgs, max_iter=3000'),
('C 후보','0.01, 0.1, 1, 10'),
('가중치 후보','무가중, 제곱근 가중, 균형 가중'),
('비교 규모','가중치 3종 × C 4종 = 12개 설정'),
('설정 선택','내부 검증 fold의 평균 Average Precision 최대'),
('임계값 선택','선택된 설정의 내부 OOF 예측에서 F2 최대')],[1.4,4.4])
p('C 범위와 반복 횟수는 초기 실행안이다. 수렴 경고가 발생하면 스케일과 수렴 상태를 확인하고 반복 횟수를 늘린다. 평균 AP가 동률이면 작은 C를 우선하고, 그다음 약한 클래스 가중치를 선택하는 규칙을 미리 고정한다. AP는 평균 정밀도이며 사다리꼴 적분 PR AUC와 혼용하지 않는다.')
h('임계값 결정 순서')
p('선택된 설정으로 내부 검증 예측을 모은다. 가능한 점수 경계에서 F2를 계산하고 최대값의 임계값을 선택한다. F2가 동률이면 Precision이 높은 임계값, 이후 더 높은 임계값을 선택한다. 선택한 설정으로 외부 학습 전체를 재학습한 뒤, 정해둔 임계값을 외부 검증에 그대로 적용한다.')
p('F2 = 5 × Precision × Recall / (4 × Precision + Recall)')
p('F2는 불량 누락을 더 중시하는 현재 목적에 맞춘 제안이다. 실제 오탐 허용량이 정해지면 그 제약 아래 Recall을 최대화하는 기준으로 바꿀 수 있다. 가중치와 임계값은 모두 검출률에 영향을 주므로, 외부 검증 결과를 보고 반복적으로 선택하지 않는다.')
h('결과 보고와 최종 적합')
p('각 외부 fold의 AP, Precision, Recall, F1, F2, TP·FP·FN·TN을 기록한다. 네 fold의 평균과 표준편차 및 전체 OOF 예측을 합친 결과를 함께 제시한다. fold별 임계값이 다르므로 전체 분류 지표는 각 행이 속한 fold의 판정을 합쳐 계산한다. 정확도는 보조 지표로 둔다.')
p('기준 비교는 무가중 모델의 임계값 0.5, 무가중 모델의 내부 선택 임계값, 가중 모델의 내부 선택 임계값으로 구성한다. 평가 종료 후에는 전체 라벨 데이터의 그룹 교차검증으로 설정과 임계값을 다시 정하고 최종 모델을 적합한다. 외부 평가 점수와 최종 모델의 학습 점수를 구분한다.')

page('5 행별 손실과 예측 확률 시각화')
p('외부 검증 예측을 원본 행 순서로 모아 모든 행에 하나의 OOF 점수를 부여한다. 각 점은 해당 행과 동일 입력 패턴을 학습에서 제외한 모델의 결과이다. 실제 학습 결과가 나오기 전에는 예시 점을 실측 결과처럼 표시하지 않는다.')
table(['항목','그래프 설계'],[
('X축','OOF 불량 예측 점수 p_i, 범위 0~1'),
('Y축','가중치 적용 전 행별 로그 손실 ℓ_i'),
('색상','실제 정상은 파랑, 실제 불량은 빨강'),
('패널','외부 fold별 4개 패널과 전체 통합 그림'),
('임계값','fold별 패널에 해당 임계값 세로 점선'),
('겹침 처리','투명도 또는 동일 좌표의 관측 개수 표시')],[1.4,4.4])
p('왼쪽 위의 빨간 점은 불량에 낮은 점수를 준 사례이고, 오른쪽 위의 파란 점은 정상에 높은 점수를 준 사례이다. 아래쪽 점은 실제 라벨에 부합하는 예측이다. 손실은 라벨과 예측 점수로 결정되므로 정상·불량 두 곡선 위에 놓이며, 이 곡선 모양 자체가 모델의 우수성을 뜻하지는 않는다.')
p('가중 손실 a_iℓ_i는 별도 그림에 표시한다. 외부 학습에서 결정된 클래스 가중치를 해당 외부 검증 행에 적용하여 비용 관점의 진단값으로 계산한다. 이는 실제 학습 중 관측된 손실이 아니며, 가중치가 다른 모델·fold의 원시 가중 손실만으로 예측 성능 순위를 정하지 않는다. L2 규제항은 행에 나누어 넣지 않는다.')
p('통합 그림에는 단일 임계값을 그리지 않는다. 정상·불량별 점수 분포와 PR 곡선을 보조 그림으로 제공한다. 원본 행 번호는 시간으로 확인되지 않았으므로 시간 축처럼 연결하지 않는다. 학습 진행에 따른 손실 감소를 보고 싶다면 별도의 반복별 학습 기록이 필요하다.')
h('저장할 결과와 해석 범위')
p('행별 결과는 source_id, pattern_id, outer_fold, y_true, p_oof, threshold, y_pred, log_loss, class_weight, weighted_log_loss로 저장한다. log 계산에서는 수치 오류 방지를 위해 확률을 작은 ε와 1−ε 사이로 제한한다. 상충 패턴 여부와 불량 전용 여부는 사후 분석용 메타데이터로만 연결한다.')
p('CN7은 불량 전용 패턴과 상충 패턴의 탐지 결과를 분리해서 보고한다. RG3는 같은 입력의 정상·불량을 함께 확인하여 누락 감소가 오탐 증가와 어떻게 연결되는지 설명한다. 전체 데이터 EDA를 이미 수행했으므로 교차검증 결과는 개발 단계의 근거로 해석하고 향후 독립 데이터에서 재확인한다.')

page('6 실행 순서와 참고 자료')
for text in [
 '1. 원본 라벨과 24개 입력을 확인하고 pattern_id를 고정한다.',
 '2. 두 데이터셋 각각 외부 4-fold를 생성하고 라벨 수와 그룹 누수를 확인한다.',
 '3. 외부 학습 구간마다 내부 2-fold를 구성한다. 전처리와 클래스 가중치는 내부 학습에서만 추정한다.',
 '4. 12개 설정을 내부 AP로 비교하고, 선택된 설정의 내부 OOF 점수로 F2 임계값을 결정한다.',
 '5. 외부 학습 전체로 재적합하여 외부 검증 점수와 판정을 저장한다.',
 '6. OOF 평가표, 행별 손실 그래프, 점수 분포와 PR 곡선을 생성하고 실패 사례를 검토한다.',
 '7. 평가를 정리한 뒤 전체 라벨 데이터로 최종 모델과 임계값을 결정한다.'
]: p(text)
h('재현성을 위한 기록')
p('데이터 파일 해시, 라이브러리 버전, 분할 seed, pattern_id와 fold 매핑, 후보 설정, 선택된 C·가중치·임계값, 수렴 여부를 저장한다. 4-fold 분할은 이후 모델 비교에도 동일하게 재사용한다. fold 수나 seed를 성능이 좋아지는 방향으로 사후 선택하지 않는다.')
h('근거 자료')
for t in [
 '프로젝트 데이터: data/origin/moldset_labeled_cn7.csv 및 moldset_labeled_rg3.csv. 행 수와 클래스 수는 문서 생성 시 재확인하였다.',
 '프로젝트 전처리: data/processed/cn7 및 rg3의 preprocessing_manifest.json과 fold_summary.csv. 현재 파일의 3-fold와 본 문서의 4-fold 설계를 구분하였다.',
 '제공 자료: 04. Guidebook_Molding-압축됨.pdf. 사전 검토에서 확인한 파일별 표준화 설명을 반영하였다.',
 'scikit-learn LogisticRegression: https://scikit-learn.org/stable/modules/generated/sklearn.linear_model.LogisticRegression.html',
 'scikit-learn StandardScaler: https://scikit-learn.org/stable/modules/generated/sklearn.preprocessing.StandardScaler.html',
 'scikit-learn StratifiedGroupKFold: https://scikit-learn.org/stable/modules/generated/sklearn.model_selection.StratifiedGroupKFold.html',
 'scikit-learn 임계값 조정: https://scikit-learn.org/stable/modules/classification_threshold.html'
]: p(t)
file=OUT/'CN7_RG3_로지스틱회귀_학습전략.docx'
doc.save(file)
check=Document(file)
assert len(check.tables)==4
assert all(str(s[1]) in ' '.join(c.text for t in check.tables for r in t.rows for c in r.cells) for s in stats)
(OUT/'verification.json').write_text(json.dumps({'document':file.name,'sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'data_counts':stats,'models_trained':False,'existing_folds_modified':False,'planned_outer_folds':4,'paragraphs':len(check.paragraphs),'tables':len(check.tables),'visual_review':False},ensure_ascii=False,indent=2),encoding='utf-8')
print(str(file))
